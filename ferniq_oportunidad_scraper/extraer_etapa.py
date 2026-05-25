import json
import os
import re
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl.styles import Alignment
from playwright.sync_api import Error, Locator, Page, TimeoutError, sync_playwright

from config_utils import load_dotenv
from guardar_sesion import get_chrome_path, save_session, try_auto_login


LOGIN_URL = "https://ferniq.fernfutures.com/"
URL = "https://ferniq.fernfutures.com/app/opportunity-flow/stage"
BASE_DIR = Path(__file__).resolve().parent
SESSION_FILE = BASE_DIR / "sesion_ferniq.json"
SESSION_STORAGE_FILE = BASE_DIR / "sesion_ferniq_session_storage.json"
EXPORTS_DIR = BASE_DIR / "exports"
OUTPUT_FILE = BASE_DIR / "resultado_etapa_ferniq.xlsx"

COMERCIALES = [
    "Ismael Serrano",
    "Javier Barcelo",
    "Jesus Cabello",
    "Miguel Angel Haro",
    "Paco Buendia",
    "Agustin Parejo",
    "Ramón Jimenez",
]

MES = "Mayo 2026"

STAGE_NAMES = [
    "Aceptado",
    "Anulado / Perdido",
    "Contacto inicial",
    "Envio de presupuesto",
    "Negociacion",
    "Oportunidad de valor",
    "Por confirmar / Sin respuesta",
    "Proceso de compras",
]

STAGE_TAB_PATTERN = re.compile(r"\bEtapa\b", re.I)
LOGIN_TEXT_PATTERNS = [
    re.compile(r"iniciar sesion", re.I),
    re.compile(r"sign in", re.I),
    re.compile(r"log in", re.I),
    re.compile(r"password", re.I),
    re.compile(r"correo", re.I),
    re.compile(r"usuario", re.I),
]
NOT_FOUND_TEXT_PATTERNS = [
    re.compile(r"404", re.I),
    re.compile(r"not found", re.I),
    re.compile(r"page not found", re.I),
]

SPINNER_SELECTORS = [
    "[aria-busy='true']",
    ".loading",
    ".loader",
    ".spinner",
    ".mat-progress-spinner",
    ".mat-mdc-progress-spinner",
]

EMPTY_COLUMNS = [
    "comercial",
    "mes",
    "etapa",
    "cantidad",
    "valor_texto",
    "precio",
    "fecha_extraccion",
    "origen",
]

COMERCIAL_ALIASES = {
    "agustin parejo": ["Mi usuario"],
    "ramon jimenez": ["Ramón Jimenez"],
}


def log(message: str) -> None:
    now = datetime.now().strftime("%H:%M:%S")
    print(f"[{now}] {message}")


def get_parallel_workers() -> int:
    raw_value = os.getenv("FERNIQ_STAGE_MAX_WORKERS", "").strip()
    default_workers = 3
    max_workers = 4

    if not raw_value:
        return default_workers

    try:
        parsed = int(raw_value)
    except ValueError:
        log(
            "Valor invalido en FERNIQ_STAGE_MAX_WORKERS. "
            f"Se usaran {default_workers} workers."
        )
        return default_workers

    return max(1, min(max_workers, parsed))


def normalize_text(value: str) -> str:
    if value is None:
        return ""
    value = unicodedata.normalize("NFKD", value)
    value = "".join(char for char in value if not unicodedata.combining(char))
    return re.sub(r"\s+", " ", value).strip().lower()


def slugify_filename(value: str) -> str:
    cleaned = re.sub(r"[^a-zA-Z0-9._-]+", "_", normalize_text(value))
    return cleaned.strip("_") or "archivo"


def first_visible(
    locators: Iterable[Locator],
    description: str,
    timeout_ms: int = 10000,
) -> Locator:
    locator = find_visible(locators, timeout_ms=timeout_ms)
    if locator is not None:
        return locator
    raise RuntimeError(f"No se encontro un locator visible para: {description}")


def find_visible(
    locators: Iterable[Locator],
    timeout_ms: int = 10000,
) -> Locator | None:
    locators = list(locators)
    deadline = time.monotonic() + (timeout_ms / 1000.0)

    while time.monotonic() < deadline:
        for locator in locators:
            try:
                if locator.count() and locator.first.is_visible():
                    return locator.first
            except Error:
                continue
        time.sleep(0.25)

    return None


def click_locator(locator: Locator, description: str) -> None:
    locator.scroll_into_view_if_needed()
    locator.click(timeout=10000)
    log(f"{description}: OK")


def wait_for_dashboard_settle(page: Page) -> None:
    try:
        page.wait_for_load_state("networkidle", timeout=10000)
    except TimeoutError:
        pass

    for selector in SPINNER_SELECTORS:
        try:
            page.locator(selector).first.wait_for(state="hidden", timeout=3000)
        except TimeoutError:
            continue
        except Error:
            continue


def locator_is_selected(locator: Locator) -> bool:
    for attribute in ("aria-selected", "aria-pressed", "data-state"):
        try:
            value = locator.get_attribute(attribute, timeout=1000)
        except Error:
            value = None
        if value and value.lower() in {"true", "active", "selected"}:
            return True
    return False


def is_stage_view_ready(page: Page) -> bool:
    current_url = page.url.lower()
    if "my-opportunities" in current_url:
        return False

    if "opportunity-flow" in current_url:
        ready_markers = [
            page.get_by_text(re.compile(r"FO\s*-\s*Etapa", re.I)),
            page.get_by_text(re.compile(r"\bEtapa\b", re.I)),
            *get_comercial_trigger_candidates(page),
        ]
        return find_visible(ready_markers, timeout_ms=2000) is not None

    stage_markers = [page.get_by_text(re.compile(re.escape(stage), re.I)) for stage in STAGE_NAMES[:3]]
    if find_visible(stage_markers, timeout_ms=1500) is not None:
        return find_visible(get_comercial_trigger_candidates(page), timeout_ms=2000) is not None

    return False


def is_login_page(page: Page) -> bool:
    candidates = [
        page.locator("input[type='password']"),
        page.locator("input[type='email']"),
        page.locator("input[name*='user' i]"),
        page.locator("input[name*='email' i]"),
    ]
    if find_visible(candidates, timeout_ms=1000) is not None:
        return True

    text_candidates = [page.get_by_text(pattern) for pattern in LOGIN_TEXT_PATTERNS]
    return find_visible(text_candidates, timeout_ms=1000) is not None


def get_comercial_trigger_candidates(page: Page) -> list[Locator]:
    return [
        page.locator("app-seller-selector ion-select"),
        page.locator("ion-select[aria-haspopup='listbox']"),
        page.locator("ion-select"),
        page.get_by_role("button", name=re.compile(r"Comercial", re.I)),
        page.get_by_role("combobox", name=re.compile(r"Comercial", re.I)),
        page.get_by_label(re.compile(r"Comercial", re.I)),
        page.get_by_text(re.compile(r"^Comercial$", re.I)),
        page.locator("[placeholder*='Comercial' i]"),
    ]


def page_debug_context(page: Page) -> str:
    try:
        title = page.title()
    except Error:
        title = ""
    return f"URL actual: {page.url}. Titulo: {title}"


def redirect_to_login(page: Page) -> None:
    try:
        page.goto(LOGIN_URL, wait_until="domcontentloaded", timeout=15000)
    except TimeoutError:
        pass
    except Error:
        pass


def try_auto_login_in_page(page: Page) -> bool:
    load_dotenv()
    email = os.getenv("FERNIQ_EMAIL", "").strip()
    password = os.getenv("FERNIQ_PASSWORD", "").strip()
    if not email or not password:
        return False

    try:
        return try_auto_login(page, email, password)
    except Exception as exc:
        log(f"Login automatico en extraccion no completado: {exc}")
        return False


def load_session_storage_payload() -> dict | None:
    if not SESSION_STORAGE_FILE.exists():
        return None

    try:
        payload = json.loads(SESSION_STORAGE_FILE.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return None

    if not isinstance(payload, dict):
        return None

    origin = payload.get("origin")
    entries = payload.get("entries")
    if not isinstance(origin, str) or not isinstance(entries, dict):
        return None

    cleaned_entries = {
        str(key): "" if value is None else str(value)
        for key, value in entries.items()
    }
    return {"origin": origin, "entries": cleaned_entries}


def restore_session_storage(context) -> None:
    payload = load_session_storage_payload()
    if not payload:
        return

    payload_json = json.dumps(payload, ensure_ascii=False)
    context.add_init_script(
        script=f"""
        (() => {{
          const payload = {payload_json};
          if (window.location.origin !== payload.origin) {{
            return;
          }}
          for (const [key, value] of Object.entries(payload.entries)) {{
            window.sessionStorage.setItem(key, value);
          }}
        }})();
        """
    )


def is_not_found_page(page: Page) -> bool:
    try:
        title = page.title()
    except Error:
        title = ""

    if re.search(r"404|not found", title, re.I):
        return True

    text_candidates = [page.get_by_text(pattern) for pattern in NOT_FOUND_TEXT_PATTERNS]
    return find_visible(text_candidates, timeout_ms=1000) is not None


def open_stage_from_home(page: Page) -> bool:
    try:
        page.goto(LOGIN_URL, wait_until="domcontentloaded", timeout=15000)
    except TimeoutError:
        pass
    except Error:
        pass

    wait_for_dashboard_settle(page)
    if is_login_page(page):
        if not try_auto_login_in_page(page):
            return False
        wait_for_dashboard_settle(page)
    if is_stage_view_ready(page):
        return True

    menu_button = find_visible(
        [
            page.locator("ion-menu-button"),
            page.get_by_role("button", name=re.compile(r"menu", re.I)),
        ],
        timeout_ms=5000,
    )
    if menu_button is not None:
        click_locator(menu_button, "Menu principal")
        page.wait_for_timeout(1000)

    candidates = [
        page.locator("ion-item").filter(has_text=re.compile(r"Flujo de oportunidad", re.I)),
        page.get_by_text(re.compile(r"Flujo de oportunidad", re.I)),
        page.get_by_role("link", name=re.compile(r"FO\s*-\s*Etapa", re.I)),
        page.get_by_role("button", name=re.compile(r"FO\s*-\s*Etapa", re.I)),
        page.get_by_text(re.compile(r"FO\s*-\s*Etapa", re.I)),
        page.get_by_role("link", name=re.compile(r"\bEtapa\b", re.I)),
        page.get_by_role("button", name=re.compile(r"\bEtapa\b", re.I)),
        page.locator("a[href*='opportunity-flow' i]"),
        page.locator("a[href*='stage' i]"),
    ]

    target = find_visible(candidates, timeout_ms=8000)
    if target is None:
        return False

    try:
        href = target.get_attribute("href", timeout=1000)
    except Error:
        href = None

    if href:
        try:
            page.goto(href, wait_until="domcontentloaded", timeout=15000)
        except TimeoutError:
            pass
        except Error:
            click_locator(target, "Acceso a FO - Etapa")
    else:
        click_locator(target, "Acceso a FO - Etapa")

    wait_for_dashboard_settle(page)
    return is_stage_view_ready(page)


def open_stage_view(page: Page) -> None:
    try:
        page.goto(URL, wait_until="domcontentloaded", timeout=15000)
    except TimeoutError:
        pass
    except Error:
        pass

    wait_for_dashboard_settle(page)

    if is_login_page(page):
        if try_auto_login_in_page(page):
            wait_for_dashboard_settle(page)
            if is_stage_view_ready(page):
                return
        else:
            redirect_to_login(page)
            raise RuntimeError(
                "La sesion no esta activa o la web ha redirigido al login. "
                f"Ejecuta guardar_sesion.py otra vez. {page_debug_context(page)}"
            )

    if is_stage_view_ready(page):
        return

    if is_not_found_page(page):
        log("La URL directa de Etapa devuelve 404; se intentara abrir desde la home")

    if open_stage_from_home(page):
        return

    if is_login_page(page):
        if not try_auto_login_in_page(page):
            redirect_to_login(page)
            raise RuntimeError(
                "La sesion no esta activa o la web ha redirigido al login. "
                f"Ejecuta guardar_sesion.py otra vez. {page_debug_context(page)}"
            )
        wait_for_dashboard_settle(page)
        if open_stage_from_home(page):
            return


def ensure_stage_tab(page: Page) -> None:
    wait_for_dashboard_settle(page)
    if is_stage_view_ready(page):
        log("Vista Etapa ya activa")
        return

    if is_login_page(page):
        redirect_to_login(page)
        raise RuntimeError(
            "La sesion no esta activa o la web ha redirigido al login. "
            f"Ejecuta guardar_sesion.py otra vez. {page_debug_context(page)}"
        )

    candidates = [
        page.get_by_role("tab", name=STAGE_TAB_PATTERN),
        page.locator("[role='tab']").filter(has_text=STAGE_TAB_PATTERN),
        page.get_by_role("button", name=STAGE_TAB_PATTERN),
        page.get_by_text(STAGE_TAB_PATTERN),
        page.locator("text=Etapa"),
    ]

    tab = find_visible(candidates, timeout_ms=5000)
    if tab is not None:
        if not locator_is_selected(tab):
            click_locator(tab, "Pestana Etapa")
        else:
            log("Pestana Etapa ya seleccionada")
        wait_for_dashboard_settle(page)
    else:
        log("Pestana Etapa no visible; se intentara continuar si Comercial ya esta disponible")

    if is_stage_view_ready(page):
        return

    if is_login_page(page):
        redirect_to_login(page)
        raise RuntimeError(
            "La sesion no esta activa o la web ha redirigido al login. "
            f"Ejecuta guardar_sesion.py otra vez. {page_debug_context(page)}"
        )

    raise RuntimeError(
        "No se detecto la vista de Etapa ni el selector Comercial. "
        f"{page_debug_context(page)}"
    )


def get_comercial_trigger(page: Page) -> Locator:
    candidates = get_comercial_trigger_candidates(page)
    return first_visible(candidates, "selector Comercial")


def canonical_comercial_name(value: str) -> str:
    normalized = normalize_text(value)
    alias_to_canonical = {
        "mi usuario": "Agustin Parejo",
        "ramon jimenez": "Ramon Jimenez",
    }
    return alias_to_canonical.get(normalized, value.strip())


def is_valid_comercial_option(value: str) -> bool:
    normalized = normalize_text(value)
    if not normalized:
        return False

    invalid_values = {
        "comercial",
        "selecciona",
        "seleccione",
        "select",
        "todos",
        "all",
    }
    return normalized not in invalid_values


def get_available_comerciales(page: Page) -> list[str]:
    candidates = [
        page.locator("app-seller-selector ion-select-option"),
        page.locator("ion-select-option"),
    ]

    detected: list[str] = []
    seen: set[str] = set()

    for locator in candidates:
        try:
            total = locator.count()
        except Error:
            continue

        for index in range(total):
            option = locator.nth(index)
            try:
                text = option.inner_text(timeout=1000).strip()
            except Error:
                try:
                    text = (option.get_attribute("value", timeout=1000) or "").strip()
                except Error:
                    text = ""

            if not is_valid_comercial_option(text):
                continue

            canonical = canonical_comercial_name(text)
            normalized = normalize_text(canonical)
            if normalized in seen:
                continue

            seen.add(normalized)
            detected.append(canonical)

        if detected:
            break

    return detected


def get_comerciales_to_process(page: Page) -> list[str]:
    detected = get_available_comerciales(page)
    if detected:
        official_by_normalized = {
            normalize_text(canonical_comercial_name(comercial)): canonical_comercial_name(comercial)
            for comercial in COMERCIALES
        }
        filtered = [
            official_by_normalized[normalize_text(canonical_comercial_name(comercial))]
            for comercial in detected
            if normalize_text(canonical_comercial_name(comercial)) in official_by_normalized
        ]
        if filtered:
            return filtered
        return detected
    return COMERCIALES


def open_comercial_popup(page: Page) -> Locator:
    trigger = get_comercial_trigger(page)
    click_locator(trigger, "Selector Comercial")

    popup_candidates = [
        page.locator("ion-alert"),
        page.locator("[role='alertdialog']"),
        page.get_by_role("dialog"),
        page.locator("[role='dialog']"),
        page.locator(".cdk-overlay-pane"),
        page.locator(".v-overlay__content"),
        page.locator(".mat-mdc-select-panel"),
        page.locator(".alert-wrapper"),
    ]

    popup = first_visible(popup_candidates, "popup de Comercial")
    popup.wait_for(state="visible", timeout=10000)
    return popup


def uncheck_all_selected(popup: Locator) -> None:
    role_checkboxes = popup.get_by_role("checkbox")
    try:
        total = role_checkboxes.count()
    except Error:
        total = 0

    for index in range(total):
        checkbox = role_checkboxes.nth(index)
        try:
            checked = checkbox.is_checked(timeout=1000)
        except Error:
            checked = checkbox.get_attribute("aria-checked") == "true"
        if checked:
            checkbox.click(timeout=5000, force=True)

    native_checked = popup.locator("input[type='checkbox']:checked")
    try:
        native_total = native_checked.count()
    except Error:
        native_total = 0

    for index in range(native_total):
        checkbox = native_checked.nth(index)
        try:
            checkbox.click(timeout=5000, force=True)
        except Error:
            try:
                checkbox.evaluate("(el) => el.click()")
            except Error:
                continue


def select_comercial(popup: Locator, comercial: str) -> None:
    candidate_names = [comercial, *COMERCIAL_ALIASES.get(normalize_text(comercial), [])]

    for candidate_name in candidate_names:
        candidates = [
            popup.get_by_role("checkbox", name=re.compile(rf"^{re.escape(candidate_name)}$", re.I)),
            popup.get_by_label(re.compile(rf"^{re.escape(candidate_name)}$", re.I)),
            popup.get_by_text(re.compile(rf"^{re.escape(candidate_name)}$", re.I)),
        ]
        try:
            option = first_visible(candidates, f"comercial {candidate_name}", timeout_ms=5000)
            click_locator(option, f"Seleccionar comercial {comercial} ({candidate_name})")
            return
        except RuntimeError:
            pass

    checkbox_like = popup.locator("label, [role='checkbox'], .mat-checkbox, .mdc-form-field")
    total = checkbox_like.count()
    target_names = {normalize_text(name) for name in candidate_names}

    for index in range(total):
        option = checkbox_like.nth(index)
        try:
            text = option.inner_text(timeout=1000)
        except Error:
            continue
        if normalize_text(text) not in target_names:
            continue
        click_locator(option, f"Seleccionar comercial {comercial}")
        return

    alias_text = ", ".join(candidate_names)
    raise RuntimeError(f"No se encontro el comercial '{comercial}' en el popup. Alias probados: {alias_text}")


def confirm_popup(popup: Locator) -> None:
    candidates = [
        popup.get_by_role("button", name=re.compile(r"^OK$", re.I)),
        popup.get_by_text(re.compile(r"^OK$", re.I)),
    ]
    ok_button = first_visible(candidates, "boton OK del popup")
    click_locator(ok_button, "Boton OK")


def select_only_comercial(page: Page, comercial: str) -> None:
    popup = open_comercial_popup(page)
    uncheck_all_selected(popup)
    select_comercial(popup, comercial)
    confirm_popup(popup)
    wait_for_dashboard_settle(page)
    log("Seleccionado correctamente")


def limpiar_precio(valor) -> float | None:
    if valor is None:
        return None

    text = str(valor).strip()
    if not text:
        return None

    text = (
        text.replace("EUR", "")
        .replace("eur", "")
        .replace("€", "")
        .replace("\u00a0", "")
        .replace(" ", "")
    )

    if not text:
        return None

    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        parts = text.split(",")
        if len(parts[-1]) in {1, 2}:
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "." in text:
        parts = text.split(".")
        if len(parts[-1]) == 3 and len(parts) > 1:
            text = text.replace(".", "")

    text = re.sub(r"[^0-9.-]", "", text)
    if not text or text in {"-", ".", "-.", ".-"}:
        return None

    try:
        return float(text)
    except ValueError:
        return None


def parse_quantity(value: str) -> int | None:
    if not value:
        return None

    cleaned = value.replace(".", "").replace(",", "")
    cleaned = re.sub(r"[^\d-]", "", cleaned)
    if not cleaned:
        return None

    try:
        return int(cleaned)
    except ValueError:
        return None


def try_extract_download(page: Page, comercial: str, extraction_date: str) -> list[dict]:
    download_button_candidates = [
        page.get_by_role("button", name=re.compile(r"Descargar", re.I)),
        page.get_by_text(re.compile(r"^Descargar$", re.I)),
    ]

    try:
        download_button = first_visible(download_button_candidates, "boton Descargar", timeout_ms=5000)
    except RuntimeError:
        log("Boton Descargar no encontrado; se usara lectura DOM.")
        return []

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    EXPORTS_DIR.mkdir(parents=True, exist_ok=True)

    try:
        with page.expect_download(timeout=15000) as download_info:
            click_locator(download_button, "Boton Descargar")
        download = download_info.value
        suggested_name = download.suggested_filename
        file_path = EXPORTS_DIR / f"{slugify_filename(comercial)}_{timestamp}_{suggested_name}"
        download.save_as(str(file_path))
        log(f"Archivo descargado: {file_path.name}")
        records = read_download_records(file_path, comercial, extraction_date)
        if records:
            log("Datos obtenidos desde descarga")
        else:
            log("La descarga no contenia datos utiles para Etapa; se usara lectura DOM.")
        return records
    except TimeoutError:
        log("La descarga no arranco a tiempo; se usara lectura DOM.")
        return []
    except Error as exc:
        log(f"Descarga no utilizable: {exc}")
        return []


def read_download_records(file_path: Path, comercial: str, extraction_date: str) -> list[dict]:
    extension = file_path.suffix.lower()
    dataframes: list[tuple[str, pd.DataFrame]] = []

    try:
        if extension == ".csv":
            dataframes.append((file_path.name, pd.read_csv(file_path)))
        elif extension in {".xlsx", ".xls"}:
            sheets = pd.read_excel(file_path, sheet_name=None)
            dataframes.extend(sheets.items())
        else:
            return []
    except Exception as exc:
        log(f"No se pudo leer el archivo descargado {file_path.name}: {exc}")
        return []

    records: list[dict] = []
    for sheet_name, dataframe in dataframes:
        normalized = normalize_download_dataframe(
            dataframe=dataframe,
            origen=f"descarga:{file_path.name}:{sheet_name}",
            comercial=comercial,
            extraction_date=extraction_date,
        )
        records.extend(normalized)

    return deduplicate_records(records)


def normalize_column_name(name: str) -> str:
    return normalize_text(str(name)).replace(" ", "_")


def match_first_column(columns: list[str], patterns: list[str]) -> str | None:
    for column in columns:
        if any(pattern in column for pattern in patterns):
            return column
    return None


def normalize_download_dataframe(
    dataframe: pd.DataFrame,
    origen: str,
    comercial: str,
    extraction_date: str,
) -> list[dict]:
    if dataframe.empty:
        return []

    df = dataframe.copy()
    df.columns = [normalize_column_name(column) for column in df.columns]
    df = df.dropna(how="all")
    if df.empty:
        return []

    columns = list(df.columns)
    etapa_col = match_first_column(columns, ["etapa", "stage", "estado"])
    cantidad_col = match_first_column(columns, ["cantidad", "count", "qty", "numero", "numero_de"])
    valor_col = match_first_column(columns, ["valor", "precio", "importe", "amount", "total"])

    records: list[dict] = []

    if etapa_col:
        for _, row in df.iterrows():
            etapa = str(row.get(etapa_col, "")).strip()
            if not etapa or etapa.lower() == "nan":
                continue

            stage_match = find_stage_name(etapa)
            if not stage_match:
                continue

            valor_texto = ""
            if valor_col:
                raw_value = row.get(valor_col)
                if pd.notna(raw_value):
                    valor_texto = str(raw_value).strip()

            cantidad = None
            if cantidad_col:
                raw_count = row.get(cantidad_col)
                if pd.notna(raw_count):
                    cantidad = parse_quantity(str(raw_count))

            records.append(
                build_record(
                    comercial=comercial,
                    etapa=stage_match,
                    cantidad=cantidad,
                    valor_texto=valor_texto,
                    precio=limpiar_precio(valor_texto),
                    extraction_date=extraction_date,
                    origen=origen,
                )
            )
        return records

    for column in df.columns:
        stage_match = find_stage_name(column)
        if not stage_match:
            continue
        series = df[column].dropna()
        if series.empty:
            continue
        first_value = series.iloc[0]
        value_text = str(first_value).strip()
        records.append(
            build_record(
                comercial=comercial,
                etapa=stage_match,
                cantidad=parse_quantity(value_text),
                valor_texto=value_text,
                precio=limpiar_precio(value_text),
                extraction_date=extraction_date,
                origen=origen,
            )
        )

    return records


def build_record(
    comercial: str,
    etapa: str,
    cantidad: int | None,
    valor_texto: str,
    precio: float | None,
    extraction_date: str,
    origen: str,
) -> dict:
    return {
        "comercial": comercial,
        "mes": MES,
        "etapa": etapa,
        "cantidad": cantidad,
        "valor_texto": valor_texto.strip(),
        "precio": precio,
        "fecha_extraccion": extraction_date,
        "origen": origen,
    }


def find_stage_name(text: str) -> str | None:
    normalized = normalize_text(text)
    for stage in STAGE_NAMES:
        if normalize_text(stage) in normalized:
            return canonical_stage_name(stage)
    return None


def canonical_stage_name(stage: str) -> str:
    mapping = {
        "envio de presupuesto": "Envio de presupuesto",
        "negociacion": "Negociacion",
    }
    normalized = normalize_text(stage)
    return mapping.get(normalized, stage)


def collect_stage_blocks(page: Page) -> list[str]:
    return page.evaluate(
        """
        ({ stageNames }) => {
          const normalize = (text) =>
            (text || "")
              .replace(/\\u00a0/g, " ")
              .replace(/\\s+/g, " ")
              .trim();

          const visible = (el) => {
            const style = window.getComputedStyle(el);
            if (style.display === "none" || style.visibility === "hidden" || Number(style.opacity) === 0) {
              return false;
            }
            const rect = el.getBoundingClientRect();
            return rect.width > 0 && rect.height > 0;
          };

          const unique = new Set();
          const results = [];
          const normalizedNames = stageNames.map((name) => normalize(name).toLowerCase());

          for (const el of document.querySelectorAll("body *")) {
            if (!visible(el)) continue;
            const text = normalize(el.innerText);
            if (!text || text.length > 240) continue;
            const lowered = text.toLowerCase();
            if (!normalizedNames.some((name) => lowered.includes(name))) continue;
            const lineCount = text.split(/\\n+/).length;
            if (lineCount > 8) continue;
            if (!unique.has(text)) {
              unique.add(text);
              results.push(text);
            }
          }

          return results;
        }
        """,
        {"stageNames": STAGE_NAMES},
    )


def extract_records_from_block(block: str, comercial: str, extraction_date: str) -> list[dict]:
    stage = find_stage_name(block)
    if not stage:
        return []

    lines = [line.strip() for line in re.split(r"[\n\r]+", block) if line.strip()]
    merged_text = " | ".join(lines)

    price_match = re.search(
        r"(€\s*[\d.,]+|[\d.,]+\s*€|EUR\s*[\d.,]+|[\d.,]+\s*EUR)",
        merged_text,
        flags=re.I,
    )
    valor_texto = price_match.group(1).strip() if price_match else ""
    precio = limpiar_precio(valor_texto)

    quantity = None
    quantity_candidates = re.findall(r"(?<![\d.,])\d{1,3}(?:[.,]\d{3})*(?![\d.,])", merged_text)
    for candidate in quantity_candidates:
        parsed = parse_quantity(candidate)
        if parsed is None:
            continue
        if precio is not None and abs(parsed - precio) < 0.0001:
            continue
        quantity = parsed
        break

    return [
        build_record(
            comercial=comercial,
            etapa=stage,
            cantidad=quantity,
            valor_texto=valor_texto or block,
            precio=precio,
            extraction_date=extraction_date,
            origen="dom",
        )
    ]


def extract_records_from_body_text(body_text: str, comercial: str, extraction_date: str) -> list[dict]:
    lines = [line.strip() for line in body_text.splitlines() if line.strip()]
    records: list[dict] = []

    for index, line in enumerate(lines):
        stage = find_stage_name(line)
        if not stage:
            continue

        window = lines[index : index + 4]
        window_text = " | ".join(window)

        price_match = re.search(
            r"(€\s*[\d.,]+|[\d.,]+\s*€|EUR\s*[\d.,]+|[\d.,]+\s*EUR)",
            window_text,
            flags=re.I,
        )
        valor_texto = price_match.group(1).strip() if price_match else line
        precio = limpiar_precio(valor_texto)

        quantity = None
        for candidate in re.findall(r"(?<![\d.,])\d{1,3}(?:[.,]\d{3})*(?![\d.,])", window_text):
            parsed = parse_quantity(candidate)
            if parsed is None:
                continue
            if precio is not None and abs(parsed - precio) < 0.0001:
                continue
            quantity = parsed
            break

        records.append(
            build_record(
                comercial=comercial,
                etapa=stage,
                cantidad=quantity,
                valor_texto=valor_texto,
                precio=precio,
                extraction_date=extraction_date,
                origen="dom_body",
            )
        )

    return deduplicate_records(records)


def deduplicate_records(records: list[dict]) -> list[dict]:
    best_by_key: dict[tuple[str, str], dict] = {}

    for record in records:
        key = (record["comercial"], record["etapa"])
        current = best_by_key.get(key)
        if current is None:
            best_by_key[key] = record
            continue

        current_score = score_record(current)
        new_score = score_record(record)
        if new_score > current_score:
            best_by_key[key] = record

    return list(best_by_key.values())


def score_record(record: dict) -> tuple[int, int, int]:
    return (
        1 if record.get("cantidad") is not None else 0,
        1 if record.get("precio") is not None else 0,
        len(str(record.get("valor_texto", ""))),
    )


def extract_stage_data_from_chart_clicks(page: Page, comercial: str, extraction_date: str) -> list[dict]:
    items = page.locator("app-opportunity-chart .list-item")
    try:
        total_items = items.count()
    except Error:
        total_items = 0

    if total_items == 0:
        return []

    records: list[dict] = []

    for index in range(total_items):
        item = items.nth(index)
        try:
            label_text = item.locator("ion-label").inner_text(timeout=2000)
        except Error:
            continue

        stage = find_stage_name(label_text)
        if not stage:
            continue

        try:
            item.click(timeout=10000)
            page.wait_for_timeout(500)
        except Error:
            continue

        total_card = page.locator("ion-card.bg-blue").first
        try:
            total_text = total_card.inner_text(timeout=3000)
        except Error:
            total_text = ""

        price_match = re.search(
            r"(€\s*[\d.,]+|[\d.,]+\s*€|EUR\s*[\d.,]+|[\d.,]+\s*EUR)",
            total_text,
            flags=re.I,
        )
        valor_texto = price_match.group(1).strip() if price_match else ""

        records.append(
            build_record(
                comercial=comercial,
                etapa=stage,
                cantidad=None,
                valor_texto=valor_texto or total_text or label_text,
                precio=limpiar_precio(valor_texto),
                extraction_date=extraction_date,
                origen="dom_chart_click",
            )
        )

    return deduplicate_records(records)


def extract_stage_data_from_dom(page: Page, comercial: str, extraction_date: str) -> list[dict]:
    records = extract_stage_data_from_chart_clicks(page, comercial, extraction_date)
    if records:
        log("Datos extraidos desde clicks en el grafico")
        return records

    blocks = collect_stage_blocks(page)
    records = []

    for block in blocks:
        records.extend(extract_records_from_block(block, comercial, extraction_date))

    records = deduplicate_records(records)
    if records:
        log("Datos extraidos desde bloques visibles del DOM")
        return records

    body_text = page.locator("body").inner_text(timeout=10000)
    records = extract_records_from_body_text(body_text, comercial, extraction_date)
    if records:
        log("Datos extraidos desde texto visible del body")
    return records


def results_to_dataframe(records: list[dict]) -> pd.DataFrame:
    return pd.DataFrame(records, columns=EMPTY_COLUMNS)


def build_summary_dataframe(records: list[dict]) -> pd.DataFrame:
    if not records:
        return pd.DataFrame(columns=["comercial", *STAGE_NAMES])

    detail_df = results_to_dataframe(records).copy()
    detail_df["precio"] = pd.to_numeric(detail_df["precio"], errors="coerce")

    summary_df = (
        detail_df.pivot_table(
            index="comercial",
            columns="etapa",
            values="precio",
            aggfunc="sum",
        )
        .reindex(columns=STAGE_NAMES)
        .fillna(0)
        .reset_index()
    )

    summary_df.columns.name = None
    return summary_df


def save_results(records: list[dict]) -> Path:
    detail_df = results_to_dataframe(records)
    summary_df = build_summary_dataframe(records)
    summary_df["total"] = 0
    centered_alignment = Alignment(horizontal="center", vertical="center")

    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
        summary_df.to_excel(writer, sheet_name="Resumen", index=False)
        detail_df.to_excel(writer, sheet_name="Detalle", index=False)

        summary_sheet = writer.sheets["Resumen"]
        summary_sheet.freeze_panes = "B2"
        summary_sheet.column_dimensions["A"].width = 24
        first_stage_column = 2
        last_stage_column = len(STAGE_NAMES) + 1
        total_column = last_stage_column + 1

        for column_index in range(2, len(summary_df.columns) + 1):
            column_letter = summary_sheet.cell(row=1, column=column_index).column_letter
            summary_sheet.column_dimensions[column_letter].width = 22
            for row_index in range(2, summary_sheet.max_row + 1):
                summary_sheet.cell(row=row_index, column=column_index).number_format = u'#,##0 €'

        for row_index in range(2, summary_sheet.max_row + 1):
            row_start = summary_sheet.cell(row=row_index, column=first_stage_column).coordinate
            row_end = summary_sheet.cell(row=row_index, column=last_stage_column).coordinate
            total_cell = summary_sheet.cell(row=row_index, column=total_column)
            total_cell.value = f"=SUM({row_start}:{row_end})"
            total_cell.number_format = u'#,##0 €'

        grand_total_row = summary_sheet.max_row + 1
        grand_total_start = summary_sheet.cell(row=2, column=total_column).coordinate
        grand_total_end = summary_sheet.cell(row=grand_total_row - 1, column=total_column).coordinate
        grand_total_cell = summary_sheet.cell(row=grand_total_row, column=total_column)
        grand_total_cell.value = f"=SUM({grand_total_start}:{grand_total_end})"
        grand_total_cell.number_format = u'#,##0 €'

        for sheet in writer.book.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    cell.alignment = centered_alignment

    log(f"Excel final generado: {OUTPUT_FILE}")
    return OUTPUT_FILE


def load_results_dataframe() -> pd.DataFrame:
    if not OUTPUT_FILE.exists():
        raise FileNotFoundError(f"No existe el Excel final: {OUTPUT_FILE}")
    workbook = pd.ExcelFile(OUTPUT_FILE)
    if "Detalle" in workbook.sheet_names:
        return pd.read_excel(OUTPUT_FILE, sheet_name="Detalle")
    return pd.read_excel(OUTPUT_FILE)


def is_auth_error(exc: Exception) -> bool:
    message = str(exc)
    return "La sesion no esta activa o la web ha redirigido al login." in message


def is_missing_session_error(exc: Exception) -> bool:
    return isinstance(exc, FileNotFoundError) and str(SESSION_FILE) in str(exc)


def try_refresh_session_automatically(headless: bool) -> bool:
    log("Intentando regenerar sesion automaticamente")
    try:
        save_session(allow_manual=False, headless=headless)
        return True
    except Exception as exc:
        log(f"No se pudo regenerar la sesion automaticamente: {exc}")
        return False


def run_extraction(headless: bool = False) -> Path:
    load_dotenv()
    if not SESSION_FILE.exists():
        raise FileNotFoundError(
            f"No existe la sesion guardada en {SESSION_FILE}. Ejecuta primero guardar_sesion.py."
        )

    EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
    extraction_date = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    results: list[dict] = []

    def build_launch_kwargs() -> dict:
        launch_kwargs = {"headless": headless}
        chrome_path = get_chrome_path()
        if chrome_path:
            launch_kwargs["executable_path"] = str(chrome_path)
        return launch_kwargs

    def detect_comerciales() -> list[str]:
        with sync_playwright() as p:
            browser = p.chromium.launch(**build_launch_kwargs())
            context = browser.new_context(storage_state=str(SESSION_FILE), accept_downloads=True)
            restore_session_storage(context)
            page = context.new_page()
            try:
                open_stage_view(page)
                ensure_stage_tab(page)
                return get_comerciales_to_process(page)
            finally:
                context.close()
                browser.close()

    def extract_for_comercial(comercial: str) -> list[dict]:
        with sync_playwright() as p:
            browser = p.chromium.launch(**build_launch_kwargs())
            context = browser.new_context(storage_state=str(SESSION_FILE), accept_downloads=True)
            restore_session_storage(context)
            page = context.new_page()
            try:
                open_stage_view(page)
                ensure_stage_tab(page)
                select_only_comercial(page, comercial)
                records = try_extract_download(page, comercial, extraction_date)
                if not records:
                    records = extract_stage_data_from_dom(page, comercial, extraction_date)
                return records
            finally:
                context.close()
                browser.close()

    comerciales = detect_comerciales()
    log(f"Comerciales detectados: {', '.join(comerciales)}")

    workers = min(get_parallel_workers(), max(1, len(comerciales)))
    log(f"Procesando Etapa con {workers} worker(s) en paralelo")

    if workers == 1:
        for comercial in comerciales:
            log(f"Procesando comercial: {comercial}")
            try:
                records = extract_for_comercial(comercial)
                if records:
                    results.extend(records)
                    log(f"Datos extraidos para {comercial}")
                else:
                    log(f"Sin datos visibles para {comercial}")
            except Exception as exc:
                log(f"Error procesando comercial {comercial}: {exc}")
                continue
    else:
        with ThreadPoolExecutor(max_workers=workers) as executor:
            future_to_comercial = {
                executor.submit(extract_for_comercial, comercial): comercial
                for comercial in comerciales
            }
            for future in as_completed(future_to_comercial):
                comercial = future_to_comercial[future]
                try:
                    records = future.result()
                    if records:
                        results.extend(records)
                        log(f"Datos extraidos para {comercial}")
                    else:
                        log(f"Sin datos visibles para {comercial}")
                except Exception as exc:
                    log(f"Error procesando comercial {comercial}: {exc}")
                    continue

    deduped = deduplicate_records(results)
    if not deduped:
        log("No se extrajeron datos; se generara un Excel vacio con columnas.")
        return save_results([])

    return save_results(deduped)


def run_extraction_with_retry() -> Path:
    try:
        log("Intentando extraccion en modo headless")
        return run_extraction(headless=True)
    except Exception as exc:
        log(f"Fallo en modo headless: {exc}")

        if is_auth_error(exc) or is_missing_session_error(exc):
            if try_refresh_session_automatically(headless=True):
                log("Sesion regenerada automaticamente; reintentando extraccion en modo headless")
                return run_extraction(headless=True)
            if try_refresh_session_automatically(headless=False):
                log("Sesion regenerada automaticamente en modo visible; reintentando extraccion en modo visible")
                return run_extraction(headless=False)

        log("Reintentando extraccion en modo visible")
        return run_extraction(headless=False)


def main() -> None:
    run_extraction(headless=False)


if __name__ == "__main__":
    main()
